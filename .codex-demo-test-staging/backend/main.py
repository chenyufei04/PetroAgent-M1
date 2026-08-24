from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
UPLOAD_DIR = PROJECT_ROOT / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
MAX_FILE_SIZE = 100 * 1024 * 1024

app = FastAPI(title="石油工程术语质控本地服务", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)


def safe_filename(filename: str | None) -> str:
    name = Path(filename or "uploaded-file").name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    return name or "uploaded-file"


def available_path(filename: str) -> Path:
    target = UPLOAD_DIR / filename
    if not target.exists():
        return target
    stem, suffix = target.stem, target.suffix
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    return UPLOAD_DIR / f"{stem}-{stamp}{suffix}"


def extract_excel_text(path: Path) -> str | None:
    if path.suffix.lower() != ".xlsx":
        return None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]
                if any(values):
                    lines.append(" ".join(values))
        return "\n".join(lines)
    finally:
        workbook.close()


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok", "upload_dir": str(UPLOAD_DIR)}


@app.post("/api/quality/upload")
async def save_quality_file(file: UploadFile = File(...)) -> dict[str, str | int | None]:
    filename = safe_filename(file.filename)
    content = await file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="文件不能超过100 MB")
    target = available_path(filename)
    target.write_bytes(content)
    try:
        extracted_text = extract_excel_text(target)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Excel解析失败：{exc}") from exc
    return {"filename": target.name, "path": str(target), "size": len(content), "extracted_text": extracted_text}
